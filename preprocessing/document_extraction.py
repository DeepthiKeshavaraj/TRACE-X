import os
import io
import re

# PDF library imports
try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

# DOCX library imports
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# Image & OCR imports
try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import pytesseract
    HAS_PYTESSERACT = True
except ImportError:
    HAS_PYTESSERACT = False

try:
    import easyocr
    _easyocr_reader = None
    HAS_EASYOCR = True
except ImportError:
    HAS_EASYOCR = False

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

SUPPORTED_TEXT_EXTS = {'.txt'}
SUPPORTED_DOC_EXTS = {'.pdf', '.docx', '.doc'}
SUPPORTED_IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}
ALLOWED_EXTENSIONS = SUPPORTED_TEXT_EXTS | SUPPORTED_DOC_EXTS | SUPPORTED_IMAGE_EXTS

def is_allowed_file(filename):
    """Checks if file extension is supported."""
    if not filename or '.' not in filename:
        return False
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS

def _ocr_image_pil(image_bytes):
    """Performs OCR on image bytes using pytesseract, easyocr, or fallback."""
    if not HAS_PIL:
        return "", "PIL (Pillow) image library is missing."

    try:
        image = Image.open(io.BytesIO(image_bytes))
        # Convert RGBA / P modes to RGB for compatibility
        if image.mode not in ('L', 'RGB'):
            image = image.convert('RGB')
    except Exception as e:
        return "", f"Could not parse image file: {str(e)}"

    # 1. Try pytesseract if available
    if HAS_PYTESSERACT:
        try:
            ocr_text = pytesseract.image_to_string(image)
            if ocr_text and ocr_text.strip():
                return ocr_text.strip(), None
        except Exception:
            pass

    # 2. Try EasyOCR if available
    if HAS_EASYOCR:
        global _easyocr_reader
        try:
            if _easyocr_reader is None:
                _easyocr_reader = easyocr.Reader(['en'], gpu=False)
            results = _easyocr_reader.readtext(image_bytes, detail=0)
            if results:
                return "\n".join(results).strip(), None
        except Exception:
            pass

    # 3. Fallback warning
    return "", "OCR engine (pytesseract or easyocr) is required to extract text from images. Please install Tesseract OCR binary or easyocr."

def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> dict:
    """
    Extracts readable text from file bytes based on file extension.
    Supports .txt, .pdf, .docx, and image formats (.png, .jpg, .jpeg, .webp, .bmp, .tif, .tiff).
    """
    if not filename or not is_allowed_file(filename):
        ext = os.path.splitext(filename or '')[1]
        return {
            'success': False,
            'text': '',
            'filename': filename or 'unknown',
            'error': f"Unsupported file format '{ext}'. Allowed formats: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        }

    if not file_bytes or len(file_bytes) == 0:
        return {
            'success': False,
            'text': '',
            'filename': filename,
            'error': "Uploaded file is empty (0 bytes)."
        }

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        return {
            'success': False,
            'text': '',
            'filename': filename,
            'error': f"File size exceeds 10 MB limit ({len(file_bytes) / (1024*1024):.2f} MB)."
        }

    ext = os.path.splitext(filename)[1].lower()

    # --- 1. Plain Text (.txt) ---
    if ext == '.txt':
        encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']
        text = None
        for enc in encodings:
            try:
                text = file_bytes.decode(enc)
                break
            except Exception:
                continue

        if text is None:
            return {
                'success': False,
                'text': '',
                'filename': filename,
                'error': "Could not decode text file using UTF-8 or standard encodings."
            }

        return {
            'success': True,
            'text': text.strip(),
            'filename': filename,
            'file_type': 'text',
            'warning': None
        }

    # --- 2. PDF Document (.pdf) ---
    elif ext == '.pdf':
        if not HAS_PYPDF:
            return {
                'success': False,
                'text': '',
                'filename': filename,
                'error': "pypdf library is not installed for PDF parsing."
            }

        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ''
                if page_text.strip():
                    extracted_pages.append(page_text.strip())

            full_text = "\n\n".join(extracted_pages).strip()

            if full_text:
                return {
                    'success': True,
                    'text': full_text,
                    'filename': filename,
                    'file_type': 'pdf',
                    'warning': None
                }
            else:
                # PDF contains no extractable text (e.g. scanned image PDF)
                ocr_text, ocr_err = _ocr_image_pil(file_bytes)
                if ocr_text:
                    return {
                        'success': True,
                        'text': ocr_text,
                        'filename': filename,
                        'file_type': 'pdf_ocr',
                        'warning': 'Extracted via OCR from scanned PDF.'
                    }
                return {
                    'success': False,
                    'text': '',
                    'filename': filename,
                    'error': f"PDF file contains scanned images or no extractable text. {ocr_err or 'OCR required.'}"
                }
        except Exception as e:
            return {
                'success': False,
                'text': '',
                'filename': filename,
                'error': f"Error parsing PDF document: {str(e)}"
            }

    # --- 3. Word Document (.docx) ---
    elif ext == '.docx':
        if not HAS_DOCX:
            return {
                'success': False,
                'text': '',
                'filename': filename,
                'error': "python-docx library is not installed for .docx parsing."
            }

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs).strip()

            if not full_text:
                return {
                    'success': False,
                    'text': '',
                    'filename': filename,
                    'error': "Word document (.docx) contains no readable text."
                }

            return {
                'success': True,
                'text': full_text,
                'filename': filename,
                'file_type': 'docx',
                'warning': None
            }
        except Exception as e:
            return {
                'success': False,
                'text': '',
                'filename': filename,
                'error': f"Error parsing Word document (.docx): {str(e)}"
            }

    elif ext == '.doc':
        return {
            'success': False,
            'text': '',
            'filename': filename,
            'error': "Legacy binary .doc format is not directly supported. Please convert file to .docx or .pdf."
        }

    # --- 4. Image Formats (.png, .jpg, .jpeg, .webp, .bmp, .tif, .tiff) ---
    elif ext in SUPPORTED_IMAGE_EXTS:
        ocr_text, ocr_err = _ocr_image_pil(file_bytes)
        if ocr_text:
            return {
                'success': True,
                'text': ocr_text,
                'filename': filename,
                'file_type': 'image_ocr',
                'warning': None
            }

        return {
            'success': False,
            'text': '',
            'filename': filename,
            'error': f"Could not extract text from image '{filename}'. {ocr_err or 'No readable text identified.'}"
        }

    return {
        'success': False,
        'text': '',
        'filename': filename,
        'error': f"Unsupported file type: {ext}"
    }


def extract_text_from_multiple_files(files_data) -> dict:
    """
    Extracts text from multiple uploaded files and combines text with file headers.
    files_data can be a list of Werkzeug FileStorage objects or tuples (file_bytes, filename).
    """
    combined_texts = []
    processed_filenames = []
    warnings = []
    errors = []

    if not files_data:
        return {
            'success': False,
            'extracted_text': '',
            'filenames': [],
            'total_files': 0,
            'warnings': ["No files were provided."]
        }

    for item in files_data:
        if hasattr(item, 'read') and hasattr(item, 'filename'):
            filename = item.filename
            file_bytes = item.read()
        elif isinstance(item, tuple) and len(item) == 2:
            file_bytes, filename = item
        else:
            continue

        if not filename:
            continue

        res = extract_text_from_file_bytes(file_bytes, filename)

        if res['success']:
            header = f"--- File: {res['filename']} ---"
            combined_texts.append(f"{header}\n{res['text']}")
            processed_filenames.append(res['filename'])
            if res.get('warning'):
                warnings.append(f"{res['filename']}: {res['warning']}")
        else:
            errors.append(f"{res['filename']}: {res.get('error', 'Failed to extract text.')}")

    if combined_texts:
        return {
            'success': True,
            'extracted_text': "\n\n".join(combined_texts),
            'filenames': processed_filenames,
            'total_files': len(processed_filenames),
            'warnings': warnings,
            'errors': errors
        }
    else:
        return {
            'success': False,
            'extracted_text': '',
            'filenames': [],
            'total_files': 0,
            'warnings': warnings,
            'errors': errors or ["Failed to extract readable text from uploaded file(s)."]
        }
