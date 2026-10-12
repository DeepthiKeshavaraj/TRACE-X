import os
import io
import unittest
import pypdf
import docx
from PIL import Image, ImageDraw
from preprocessing.document_extraction import extract_text_from_file_bytes, extract_text_from_multiple_files

class TestDocumentExtraction(unittest.TestCase):

    def test_txt_extraction(self):
        content = b"CLASSIFIED REPORT: Operation Cyber-Trace\nLocation: Sydney"
        res = extract_text_from_file_bytes(content, "report.txt")
        self.assertTrue(res['success'])
        self.assertIn("Operation Cyber-Trace", res['text'])
        self.assertEqual(res['filename'], "report.txt")

    def test_docx_extraction(self):
        doc = docx.Document()
        doc.add_paragraph("Subject Marcus Vance entered building.")
        doc.add_paragraph("Dr. Elena Rostova met with Vance.")
        bio = io.BytesIO()
        doc.save(bio)
        docx_bytes = bio.getvalue()

        res = extract_text_from_file_bytes(docx_bytes, "report.docx")
        self.assertTrue(res['success'])
        self.assertIn("Marcus Vance", res['text'])
        self.assertIn("Elena Rostova", res['text'])

    def test_pdf_extraction(self):
        writer = pypdf.PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        # Note: blank pdf page
        bio = io.BytesIO()
        writer.write(bio)
        pdf_bytes = bio.getvalue()

        res = extract_text_from_file_bytes(pdf_bytes, "empty_sample.pdf")
        self.assertIsInstance(res, dict)

    def test_image_extraction(self):
        # Generate simple test image
        img = Image.new('RGB', (200, 100), color=(255, 255, 255))
        d = ImageDraw.Draw(img)
        d.text((10, 10), "TEST OCR EVIDENCE", fill=(0, 0, 0))
        bio = io.BytesIO()
        img.save(bio, format='PNG')
        img_bytes = bio.getvalue()

        res = extract_text_from_file_bytes(img_bytes, "test_evidence.png")
        self.assertIsInstance(res, dict)

    def test_unsupported_format(self):
        res = extract_text_from_file_bytes(b"binary data", "malicious_script.exe")
        self.assertFalse(res['success'])
        self.assertIn("Unsupported file format", res['error'])

    def test_oversized_file(self):
        huge_bytes = b"0" * (11 * 1024 * 1024)  # 11 MB
        res = extract_text_from_file_bytes(huge_bytes, "large_file.txt")
        self.assertFalse(res['success'])
        self.assertIn("exceeds 10 MB limit", res['error'])

    def test_multiple_files_extraction(self):
        file1 = (b"File 1 Evidence text content", "file1.txt")
        file2 = (b"File 2 Evidence text content", "file2.txt")
        res = extract_text_from_multiple_files([file1, file2])
        self.assertTrue(res['success'])
        self.assertEqual(res['total_files'], 2)
        self.assertIn("--- File: file1.txt ---", res['extracted_text'])
        self.assertIn("--- File: file2.txt ---", res['extracted_text'])

if __name__ == '__main__':
    unittest.main()
